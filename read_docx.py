import docx
import sys

def extract_text(doc_path, txt_path):
    doc = docx.Document(doc_path)
    with open(txt_path, 'w', encoding='utf-8') as f:
        f.write("Paragraphs:\n")
        f.write("===========\n")
        for p in doc.paragraphs:
            f.write(p.text + '\n')
        
        f.write("\n\nTables:\n")
        f.write("===========\n")
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    row_data.append(cell.text.replace('\n', ' '))
                f.write(' | '.join(row_data) + '\n')
            f.write("-" * 40 + "\n")

if __name__ == "__main__":
    extract_text('System Analytics Report on IBEX System.docx', 'extracted_requirements.txt')
